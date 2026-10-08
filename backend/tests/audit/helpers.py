"""Shared helpers for the audit tests: identities, publishing, reading output."""

from __future__ import annotations

import io
import time
from pathlib import Path

import jwt
from docx import Document

JWT_SECRET = "audit-jwt-secret-that-is-at-least-32-bytes"
USER_A = "aaaaaaaa-0000-4000-8000-00000000000a"
USER_B = "bbbbbbbb-0000-4000-8000-00000000000b"


def token_for(uid: str, **claims) -> str:
    payload = {
        "sub": uid,
        "email": f"{uid[:4]}@example.com",
        "aud": "authenticated",
        "exp": int(time.time()) + 3600,
    }
    payload.update(claims)
    return jwt.encode(payload, JWT_SECRET, algorithm="HS256")


def docx_text(data: bytes) -> str:
    """Everything a reader would see: body paragraphs, table cells, headers, footers."""
    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    for section in doc.sections:
        for hf in (section.header, section.footer):
            parts.extend(p.text for p in hf.paragraphs)
    return "\n".join(parts)


def analyze(db, settings, docs: list[Path], owner_id: str, mode: str = "smart"):
    """Run a real analysis over sample documents and hand the job to ``owner_id``."""
    from docforge.document_ingest import store_source_document
    from docforge.services import analyze_documents

    sources = [
        store_source_document(db, p.name, Path(p).read_bytes(), owner_id=owner_id) for p in docs
    ]
    job = analyze_documents(db, sources, settings=settings, mode=mode)
    job.owner_id = owner_id
    db.commit()
    return job


def publish(db, settings, docs: list[Path], owner_id: str, name: str = "Template", **kw):
    from docforge.services import publish_template
    from docforge.template_registry import TemplateRegistry

    job = analyze(db, settings, docs, owner_id, mode=kw.pop("mode", "smart"))
    template, version = publish_template(
        db, job, name=name, settings=settings,
        registry=TemplateRegistry(settings.templates_dir), owner_id=owner_id, **kw,
    )
    db.commit()
    return template, version, job


def generate_bytes(db, settings, template, data: dict, **kw) -> bytes:
    """Generate through the service layer and return the stored DOCX bytes."""
    from docforge.schemas.enums import GenerationMode
    from docforge.schemas.generation import GenerationInput
    from docforge.services.generation import generate_document
    from docforge.storage import get_storage

    gen = generate_document(
        db, template,
        GenerationInput(mode=GenerationMode.STRUCTURED_JSON, data=data, **kw),
        settings=settings,
    )
    return get_storage().get_bytes(gen.output_path)


def first_dynamic_paragraph(job) -> dict:
    """The first classification that templatizes a body paragraph (not a table)."""
    for c in job.classification["classifications"]:
        if c["classification"].startswith("DYNAMIC_") and c.get("field_name"):
            return c
    raise AssertionError("sample has no dynamic paragraph")
