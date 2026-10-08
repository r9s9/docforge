"""Regression tests for four critical findings, now fixed (0.16.1):

  AUD-001  user-supplied label text (static_prefix) executed as Jinja
  AUD-002  split-run braces in an uploaded DOCX executed as Jinja
  AUD-003  publishing a version into another user's template
  AUD-004  a ".." upload reference escaping the caller's staging area

Each test asserts the behaviour the business rule requires; a failure here means
one of these vulnerabilities has regressed.
"""

from __future__ import annotations

import copy

from docx import Document

from tests.audit.helpers import (
    USER_A,
    USER_B,
    analyze,
    docx_text,
    first_dynamic_paragraph,
    generate_bytes,
    publish,
)

CANARY = "canary-4d1f-audit"
# 7 * 191 == 1337: a number no sample document contains, so seeing it proves
# the expression was evaluated rather than copied.
PAYLOAD = "{{ 7 * 191 }}|{{ cycler.__init__.__globals__.os.environ.get('DOCFORGE_AUDIT_CANARY') }}|"


def test_aud001_label_text_from_the_client_is_never_executed(http, audit_db, settings_tmp, project_docs, monkeypatch):
    """A signed-in user controls the label text kept in front of a field.

    Rule: whatever text a template contains is text. It must never be executed
    on the server, and it must never be able to read the server's environment.
    """
    monkeypatch.setenv("DOCFORGE_AUDIT_CANARY", CANARY)
    job = analyze(audit_db, settings_tmp, project_docs, USER_A)
    classifications = copy.deepcopy(job.classification["classifications"])
    target = next(
        c for c in classifications
        if c["node_id"] == first_dynamic_paragraph(job)["node_id"]
    )
    target["static_prefix"] = PAYLOAD

    r = http.post(
        "/api/templates",
        json={"analysis_job_id": job.id, "name": "Injected", "classifications": classifications},
        headers=http.as_user(USER_A),
    )
    assert r.status_code == 201, r.text
    tid = r.json()["template"]["id"]

    g = http.post(
        f"/api/templates/{tid}/generate",
        json={"mode": "structured_json", "data": {target["field_name"]: "value"}},
        headers=http.as_user(USER_A),
    )
    assert g.status_code == 200, g.text
    out = http.get(f"/api/generations/{g.json()['id']}/download", headers=http.as_user(USER_A))
    text = docx_text(out.content)

    assert "1337" not in text, "a Jinja expression in label text was evaluated on the server"
    assert CANARY not in text, "label text read the server's environment variables"


def _split_run_docx(path):
    """A document whose body contains '{{ 7 * 191 }}' split across three runs.

    This is what Word produces when an author types braces with autocorrect,
    spell-check or formatting changes in between: the characters are adjacent
    on the page but live in separate <w:t> elements.
    """
    doc = Document()
    doc.add_heading("Quarterly Report", level=1)
    p = doc.add_paragraph()
    p.add_run("Reference ")
    p.add_run("{")
    p.add_run("{ 7 * 191 }")
    p.add_run("}")
    p.add_run(" end")
    doc.add_paragraph("Body text that describes the quarter.")
    doc.save(path)
    return path


def test_aud002_braces_typed_in_an_uploaded_document_stay_literal(audit_db, settings_tmp, tmp_path):
    """Rule: text that was literally in the source document is reproduced literally."""
    src = _split_run_docx(tmp_path / "braces.docx")
    job = analyze(audit_db, settings_tmp, [src], USER_A)
    # Keep every element as fixed boilerplate, so the only way '1337' can
    # appear is the builder handing the author's literal braces to Jinja.
    result = copy.deepcopy(job.classification)
    for c in result["classifications"]:
        c["classification"] = "FIXED"
        c["field_name"] = None
    job.classification = result
    audit_db.commit()

    from docforge.services import publish_template
    from docforge.template_registry import TemplateRegistry

    template, _ = publish_template(
        audit_db, job, name="Braces", settings=settings_tmp,
        registry=TemplateRegistry(settings_tmp.templates_dir), owner_id=USER_A,
    )
    audit_db.commit()
    text = docx_text(generate_bytes(audit_db, settings_tmp, template, {}))

    assert "1337" not in text, "literal braces from the source document were executed"
    assert "{{ 7 * 191 }}" in text.replace("\u200b", ""), "the author's literal text was lost"


def test_aud003_a_user_cannot_publish_into_someone_elses_template(http, audit_db, settings_tmp, project_docs, invoice_docs):
    """Rule: only the owner of a template can add versions to it."""
    victim, _, _ = publish(audit_db, settings_tmp, project_docs, USER_A, name="A's report")
    attacker_job = analyze(audit_db, settings_tmp, invoice_docs, USER_B)

    r = http.post(
        "/api/templates",
        json={"analysis_job_id": attacker_job.id, "template_id": victim.id},
        headers=http.as_user(USER_B),
    )
    audit_db.expire_all()
    from docforge.db.models import Template

    after = audit_db.get(Template, victim.id)
    assert r.status_code in (403, 404), f"cross-tenant publish accepted: {r.status_code} {r.text[:200]}"
    assert after.latest_version == 1, "another user's template gained a version"


def test_aud004_an_upload_reference_cannot_reach_outside_the_callers_staging_area(http, audit_db, settings_tmp, project_docs):
    """Rule: a user can only hand the server files they uploaded themselves."""
    from docforge.storage import get_storage
    from docforge.template_registry import TemplateRegistry

    victim, _, _ = publish(audit_db, settings_tmp, project_docs, USER_A, name="A's report")
    victim_key = TemplateRegistry(settings_tmp.templates_dir).template_docx_key(victim.id, 1)
    assert get_storage().exists(victim_key)

    sneaky = f"uploads/incoming/{USER_B}/../../../{victim_key}"
    r = http.post(
        "/api/templates/analyze-refs",
        json={"sources": [{"key": sneaky, "filename": "mine.docx"}], "mode": "smart"},
        headers=http.as_user(USER_B),
    )

    still_there = get_storage().exists(victim_key)
    assert r.status_code == 400, f"traversal key accepted: {r.status_code} {r.text[:200]}"
    assert still_there, "the victim's published template file was deleted"
