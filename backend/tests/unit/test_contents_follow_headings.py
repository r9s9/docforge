"""A table of contents follows its headings: each entry takes its heading's tag.

Word writes each contents entry as a copy of a heading's text, linked to a
bookmark on that heading. In a template the entry is not content of its own,
so it must not become a field of its own either: it shows the heading's tag,
and a generated document's contents read like its headings.
"""

from __future__ import annotations

import io
import zipfile
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docforge.ai_classifier import classify, derive_field_definitions
from docforge.ai_classifier.contents import entry_title, is_entry_number
from docforge.assembler import assemble
from docforge.structure_normalizer import build_extraction
from docforge.template_builder import build_template_docx

HEADINGS = ["Project Overview", "Delivery Plan", "Risks and Issues"]


def _run(p, text=None, *, fld=None, instr=None, tab=False):
    r = OxmlElement("w:r")
    if fld:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), fld)
        r.append(fc)
    if instr:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r.append(it)
    if tab:
        r.append(OxmlElement("w:tab"))
    if text is not None:
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
    p.append(r)
    return r


def _contents_doc(tmp_path, *, hyperlinks=True):
    """Contents list (a real TOC field) + three bookmarked headings."""
    doc = Document()
    doc.add_paragraph("CONTENTS")
    for i, title in enumerate(HEADINGS):
        para = doc.add_paragraph(style=doc.styles["toc 1"] if "toc 1" in doc.styles else None)
        if "toc 1" not in [s.name for s in doc.styles]:
            para.style = doc.styles.add_style("toc 1", 1)
        p = para._p
        if i == 0:
            _run(p, fld="begin")
            _run(p, instr=' TOC \\o "1-3" \\h \\z \\u ')
            _run(p, fld="separate")
        host = p
        if hyperlinks:
            host = OxmlElement("w:hyperlink")
            host.set(qn("w:anchor"), f"_Toc{i}")
            p.append(host)
        _run(host, f"{i + 1}")
        _run(host, tab=True)
        _run(host, title)
        _run(host, tab=True)
        _run(host, fld="begin")
        _run(host, instr=f" PAGEREF _Toc{i} \\h ")
        _run(host, fld="separate")
        _run(host, str(i + 2))
        _run(host, fld="end")
        if i == len(HEADINGS) - 1:
            _run(p, fld="end")
    for i, title in enumerate(HEADINGS):
        h = doc.add_heading(title, level=1)
        bs = OxmlElement("w:bookmarkStart")
        bs.set(qn("w:id"), str(i))
        bs.set(qn("w:name"), f"_Toc{i}")
        be = OxmlElement("w:bookmarkEnd")
        be.set(qn("w:id"), str(i))
        h._p.insert(1, bs)
        h._p.append(be)
        doc.add_paragraph(f"Body text for {title.lower()} goes here in full sentences.")
    path = tmp_path / "contents.docx"
    doc.save(str(path))
    return path


def _build(path, mode="tags_only"):
    ext = build_extraction(path, "d0")
    result = classify(ext, None, mode=mode)
    fields = derive_field_definitions(ext, result)
    return ext, result, fields, build_template_docx(path, result, fields)


def _contents_lines(data: bytes) -> list[str]:
    d = Document(BytesIO(data))
    return [p.text for p in d.paragraphs if (p.style.name or "").lower().startswith("toc")]


@pytest.mark.parametrize("hyperlinks", [True, False])
def test_each_entry_takes_its_headings_tag(tmp_path, hyperlinks):
    ext, result, fields, tpl = _build(_contents_doc(tmp_path, hyperlinks=hyperlinks))

    # No field for the contents list itself: its entries follow the headings.
    entries = [e for e in ext.top_level_elements() if (e.style_name or "").startswith("toc")]
    by_node = {c.node_id: c for c in result.classifications}
    assert all(by_node[e.node_id].links_to for e in entries)
    assert not any(e.node_id in f.node_ids for f in fields for e in entries)

    heading_tags = [
        p.text for p in Document(BytesIO(tpl)).paragraphs if p.style.name.startswith("Heading")
    ]
    lines = _contents_lines(tpl)
    assert len(lines) == 3
    for line, tag in zip(lines, heading_tags, strict=True):
        assert tag in line, (line, tag)  # the heading's own tag
    assert lines[0].startswith("1\t") and lines[0].endswith("\t2")  # number + page kept


def test_the_contents_field_survives_for_word_to_refresh(tmp_path):
    *_, tpl = _build(_contents_doc(tmp_path))
    xml = zipfile.ZipFile(io.BytesIO(tpl)).read("word/document.xml").decode()
    assert "TOC \\o" in xml and xml.count("PAGEREF") == 3
    kinds = [k for k in ("begin", "separate", "end") for _ in range(xml.count(f'w:fldCharType="{k}"'))]
    assert kinds.count("begin") == kinds.count("end") == 4


def test_a_generated_document_lists_its_own_headings(tmp_path):
    _, _, fields, tpl = _build(_contents_doc(tmp_path))
    titles = [f for f in fields if f.field_name.endswith("_title")]
    assert len(titles) == 3
    ctx = {f.field_name: "" for f in fields}
    ctx[titles[0].field_name] = "Scope"
    ctx[titles[2].field_name] = "Open Risks"  # the middle heading is left empty
    out = assemble(tpl, ctx, fields)
    lines = _contents_lines(out)
    assert [ln.split("\t")[1] for ln in lines] == ["Scope", "Open Risks"]
    # The entry that lost its heading went with it, and every field still closes.
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    assert xml.count('w:fldCharType="begin"') == xml.count('w:fldCharType="end"')


def test_smart_mode_keeps_fixed_headings_text_in_the_contents(tmp_path):
    _, result, fields, tpl = _build(_contents_doc(tmp_path), mode="smart")
    lines = _contents_lines(tpl)
    assert [ln.split("\t")[1] for ln in lines] == HEADINGS


def test_entry_title_and_numbers():
    assert entry_title("1.1\tAgenda & Topics\t3") == "Agenda & Topics"
    assert entry_title("Annex\t12") == "Annex"
    assert entry_title("Introduction") == "Introduction"
    assert is_entry_number("1.1.2") and is_entry_number("A.") and not is_entry_number("Agenda")
