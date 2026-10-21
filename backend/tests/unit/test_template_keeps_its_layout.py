"""A template keeps the document's layout: its rows, its looks, its headings.

Found on a project plan whose template no longer read like the document: a
value cell repeated its row's label ("Document ID«Document ID»"), the eight
rows of a "label | value" information table collapsed into one looped row, a
cover's title and subtitle (and its background picture) became one repeating
list, and leaving out one section would also have removed the next chapter's
heading.
"""

from __future__ import annotations

import io
import zipfile
from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

from docforge.ai_classifier import classify, derive_field_definitions
from docforge.ai_classifier.heuristic import looks_like_data_table, looks_like_key_value_table
from docforge.schemas.extraction import TableStructure
from docforge.structure_normalizer import build_extraction
from docforge.template_builder import build_template_docx
from docforge.template_builder.builder import _real_affixes

INFO = [
    ("Document title", "Project HELIX: Enterprise Document Management Modernisation"),
    ("Document type", "Project Management Plan"),
    ("Document ID", "VLT-PMP-2026-014"),
    ("Version / Status", "1.2 Approved"),
    ("Classification", "Confidential"),
    ("Document owner", "Lukas Brandt, Programme Manager"),
]


def _ts(rows):
    return TableStructure(n_rows=len(rows), n_cols=len(rows[0]), headers=list(rows[0]), rows=rows)


# --- Label text around a tag must be the paragraph's own ------------------

def test_a_prefix_the_paragraph_does_not_have_is_dropped():
    assert _real_affixes("VLT-PMP-2026-014", "Document ID", "") == ("", "")
    assert _real_affixes("Invoice Number: INV-1", "Invoice Number: ", "") == ("Invoice Number: ", "")
    assert _real_affixes("Total (EUR)", "Total (", ")") == ("Total (", ")")
    assert _real_affixes("Ends here.", "", "nowhere") == ("", "")


def _info_doc(tmp_path):
    doc = Document()
    doc.add_heading("Document Control", level=1)
    table = doc.add_table(rows=len(INFO), cols=2)
    for row, (label, value) in zip(table.rows, INFO, strict=True):
        row.cells[0].text, row.cells[1].text = label, value
    doc.add_paragraph("This document is the property of its author and may not be copied.")
    path = tmp_path / "info.docx"
    doc.save(str(path))
    return path


def _build(path, mode="tags_only"):
    ext = build_extraction(path, "d0")
    result = classify(ext, None, mode=mode)
    return ext, result


def _table_texts(data: bytes) -> list[list[str]]:
    table = Document(BytesIO(data)).tables[0]
    return [[c.text for c in row.cells] for row in table.rows]


def test_a_value_cell_never_carries_its_rows_label(tmp_path):
    # One row: a layout table, so each cell is a field of its own.
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Document ID", "VLT-PMP-2026-014"
    doc.add_paragraph("This document is the property of its author and may not be copied.")
    path = tmp_path / "id.docx"
    doc.save(str(path))
    ext, result = _build(path)
    by_node = {c.node_id: c for c in result.classifications}
    # What a model did: the row's label handed to the value cell as a prefix.
    value_node = next(
        e.node_id for e in ext.elements
        if e.parent_node_id and (e.text or "").strip() == "VLT-PMP-2026-014"
    )
    by_node[value_node].static_prefix = "Document ID"
    fields = derive_field_definitions(ext, result)
    rows = _table_texts(build_template_docx(path, result, fields))

    value = rows[0][1]
    assert "Document ID" not in value and value.startswith("{{") and value.endswith("}}"), rows


# --- A "label | value" table keeps every row --------------------------------

def test_a_label_value_table_is_a_form_not_a_list():
    assert looks_like_key_value_table(_ts([list(r) for r in INFO]))
    assert not looks_like_data_table(_ts([list(r) for r in INFO]))
    # Real headings over a list still make a list.
    glossary = [["TERM", "DEFINITION"], ["Content type", "A reusable set of metadata"],
                ["Retention label", "A rule for how long content is kept"]]
    assert not looks_like_key_value_table(_ts(glossary))
    assert looks_like_data_table(_ts(glossary))


def test_every_row_of_an_information_table_survives(tmp_path):
    path = _info_doc(tmp_path)
    ext, result = _build(path)
    fields = derive_field_definitions(ext, result)
    data = build_template_docx(path, result, fields)

    rows = _table_texts(data)
    assert len(rows) == len(INFO)
    assert all(cell.startswith("{{") for row in rows for cell in row)
    assert "{%tr" not in zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()


# --- Only paragraphs that look alike share one repeating field ----------------

def _cover_doc(tmp_path):
    doc = Document()
    for text, size, color in (
        ("Project HELIX", 34, RGBColor(0x11, 0x22, 0x44)),
        ("Enterprise Document Management Modernisation", 15, RGBColor(0x66, 0x77, 0x88)),
    ):
        run = doc.add_paragraph().add_run(text)
        run.font.size, run.font.color.rgb = Pt(size), color
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("The project replaces the file shares with one governed platform.")
    doc.add_paragraph("It is delivered in three migration waves over twelve months.")
    path = tmp_path / "cover.docx"
    doc.save(str(path))
    return path


def test_a_title_and_its_subtitle_keep_their_own_fields(tmp_path):
    path = _cover_doc(tmp_path)
    ext, result = _build(path)
    by_node = {c.node_id: c for c in result.classifications}
    top = ext.top_level_elements()
    title, subtitle = by_node[top[0].node_id], by_node[top[1].node_id]
    assert title.field_name and subtitle.field_name
    assert title.field_name != subtitle.field_name  # one look each

    body = [by_node[e.node_id] for e in top if e.text.startswith(("The project", "It is"))]
    assert body[0].field_name == body[1].field_name  # body text still groups


def test_a_paragraph_with_a_picture_stands_alone(tmp_path):
    from docforge.ai_classifier.heuristic import holds_picture, paragraph_look

    ext = build_extraction(_cover_doc(tmp_path), "d0")
    first = ext.top_level_elements()[0]
    assert not holds_picture(first)
    first.semantic_hints.append("image")
    assert holds_picture(first)
    assert paragraph_look(first) != paragraph_look(ext.top_level_elements()[1])


# --- Leaving a section out keeps the next chapter's heading -----------------

def test_a_hidden_section_keeps_the_next_chapter_heading(tmp_path):
    doc = Document()
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("The project is approved to proceed on the basis of this plan.")
    doc.add_heading("Project Overview", level=1)
    doc.add_heading("Background", level=2)
    doc.add_paragraph("Kestrel operates road, rail and sea freight across Europe.")
    path = tmp_path / "chapters.docx"
    doc.save(str(path))
    ext, result = _build(path)
    fields = derive_field_definitions(ext, result)
    xml = zipfile.ZipFile(io.BytesIO(build_template_docx(path, result, fields))).read(
        "word/document.xml"
    ).decode()

    body = Document(BytesIO(build_template_docx(path, result, fields))).paragraphs
    texts = [p.text for p in body]
    overview = next(i for i, t in enumerate(texts) if "project_overview" in t)
    # Whatever wraps the Executive Summary closes before the next chapter.
    assert "{%p endif %}" not in texts[overview:overview + 1]
    opened = sum(t.startswith("{%p if") for t in texts[:overview])
    closed = sum(t == "{%p endif %}" for t in texts[:overview])
    assert opened == closed, texts
    assert xml.count("{%p if") == xml.count("{%p endif %}")
