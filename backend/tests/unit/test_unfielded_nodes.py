"""A placeholder must name a field the render context will actually hold.

The failure this guards against: the classifier named a blank spacing paragraph
("spacing_placeholder"), ``derive_field_definitions`` dropped it — a blank line
is spacing, not content — and the builder wrote the tag anyway from the
classification's name. Nothing could ever fill it, and the sample-filled preview
died on it with "UndefinedError: 'spacing_placeholder' is undefined".

Two things had to be true for that to happen, so both are tested: the builder
must not invent the tag, and a template that already carries one (published
before the fix) must still render.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from docforge.ai_classifier import derive_field_definitions
from docforge.assembler import assemble
from docforge.schemas.classification import ClassificationResult, ElementClassification
from docforge.schemas.enums import ClassificationType, FieldType
from docforge.services.preview_docx import _sample_context
from docforge.structure_normalizer import build_extraction
from docforge.template_builder.builder import build_template_docx


def _classify(node_id: str, field_name: str, kind=ClassificationType.DYNAMIC_TEXT):
    return ElementClassification(
        node_id=node_id,
        classification=kind,
        field_name=field_name,
        field_type=FieldType.TEXT,
        confidence=0.8,
        rationale="test",
    )


@pytest.fixture
def named_blank_line(tmp_path):
    """A document whose blank spacing paragraph the classifier gave a name."""
    doc = Document()
    doc.add_paragraph("Acme Corporation")
    doc.add_paragraph("")  # spacing — named by the AI, dropped as a field
    doc.add_paragraph("Prepared by the delivery team.")
    src = tmp_path / "letter.docx"
    doc.save(str(src))

    extraction = build_extraction(str(src), "letter")
    blank = next(e for e in extraction.elements if not (e.text or "").strip())
    result = ClassificationResult(
        extraction_document_id="letter",
        classifications=[
            _classify(extraction.elements[0].node_id, "company_name"),
            _classify(blank.node_id, "spacing_placeholder"),
        ],
        sections=[],
        document_type_guess="letter",
    )
    fields = derive_field_definitions(extraction, result)
    return str(src), result, fields


def test_no_field_means_no_placeholder(named_blank_line):
    src, result, fields = named_blank_line
    assert [f.field_name for f in fields] == ["company_name"], "the blank line became a field"

    built = Document(BytesIO(build_template_docx(src, result, fields)))
    texts = [p.text for p in built.paragraphs]

    assert "{{ company_name }}" in texts, "the real field lost its placeholder"
    assert not any("spacing_placeholder" in t for t in texts), (
        "the builder invented a placeholder for a field nobody derived"
    )


def test_the_sample_filled_preview_builds(named_blank_line):
    """What the review screen does: build, then render with «Label» samples."""
    src, result, fields = named_blank_line
    template = build_template_docx(src, result, fields)

    rendered = Document(BytesIO(assemble(template, _sample_context(fields), fields)))

    assert "«Company Name»" in [p.text for p in rendered.paragraphs]


def test_a_stray_tag_costs_only_its_own_spot(tmp_path):
    """Templates published before the fix still carry the tag — render anyway.

    Escaping probes every value for ``__html__``; an undefined that answered
    that probe got called and raised, so one dead tag failed the whole document
    rather than the one paragraph it sat in.
    """
    doc = Document()
    doc.add_paragraph("{{ company_name }}")
    doc.add_paragraph("{{ spacing_placeholder }}")
    doc.add_paragraph("Prepared by the delivery team.")
    stray = tmp_path / "stale_template.docx"
    doc.save(str(stray))

    out = assemble(stray.read_bytes(), {"company_name": "Acme & Sons"}, [])

    texts = [p.text for p in Document(BytesIO(out)).paragraphs]
    assert texts == ["Acme & Sons", "", "Prepared by the delivery team."]
