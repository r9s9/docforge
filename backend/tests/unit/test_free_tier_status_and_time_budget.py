"""Free-tier users see the model that serves them; serverless requests finish in time.

Two faults reported together. After spending a first free AI action, a user's
status read "openai/gpt-4o-mini": counting the action created an empty settings
row, and the page filled its blanks with old OpenAI defaults while reporting
the free tier as active. And on Vercel an analysis made a dozen or more model
calls in a row inside one request, ran past the platform's limit, and the
browser saw only "network error". Independent calls now run together, and the
AI work inside a serverless request stops at a time budget so the request
still returns a finished result.
"""

from __future__ import annotations

import contextvars
import threading
import time

import pytest
from docx import Document
from fastapi.testclient import TestClient

from docforge import ai_quota
from docforge.ai import deadline
from docforge.ai.client import LLMClient, LLMError, LLMUnavailable
from docforge.ai.errors import LLMCancelled
from docforge.ai.parallel import run_all, slots_for
from docforge.config import get_settings
from docforge.db.models import UserAIConfig
from docforge.settings_store import AIConfig


@pytest.fixture
def free_settings(monkeypatch):
    s = get_settings()
    monkeypatch.setattr(s, "free_ai_enabled", True)
    monkeypatch.setattr(s, "free_ai_api_key", "sk-ant-test")
    monkeypatch.setattr(s, "free_ai_provider", "anthropic")
    monkeypatch.setattr(s, "free_ai_base_url", "https://api.anthropic.com")
    monkeypatch.setattr(s, "free_ai_model", "claude-haiku-5-5")
    monkeypatch.setattr(s, "free_ai_limit", 10)
    return s


# --- What the status shows -------------------------------------------------

def test_a_free_user_still_sees_haiku_after_spending_an_action(monkeypatch, free_settings):
    from docforge.api.routes.settings import _ai_dto

    row = UserAIConfig(owner_id="u1", free_used=1)  # what spending a credit leaves
    monkeypatch.setattr(ai_quota, "_row", lambda oid: row)

    ai = _ai_dto(row, "u1")

    assert ai["source"] == "free" and ai["active"] is True
    assert (ai["serving_provider"], ai["serving_model"]) == ("anthropic", "claude-haiku-5-5")
    # The form offers the recommended default, not the old OpenAI one.
    assert (ai["provider"], ai["model"]) == ("anthropic", "claude-haiku-5-5")
    assert "gpt-4o-mini" not in str(ai)


def test_an_own_key_user_sees_their_own_model(monkeypatch, free_settings):
    from docforge.api.routes.settings import _ai_dto

    row = UserAIConfig(
        owner_id="u2", provider="openai", base_url="https://openrouter.ai/api/v1",
        api_key="sk-or-v1-mine", model="vendor/model-x", enabled=True,
    )
    monkeypatch.setattr(ai_quota, "_row", lambda oid: row)

    ai = _ai_dto(row, "u2")

    assert ai["source"] == "own"
    assert ai["serving_model"] == "vendor/model-x" == ai["model"]


# --- The time budget --------------------------------------------------------

def test_no_budget_changes_nothing():
    assert deadline.remaining() is None
    assert deadline.budgeted(240) == 240
    deadline.check()


def test_a_call_gets_only_the_time_that_is_left():
    with deadline.time_budget(60):
        assert deadline.budgeted(240) <= 60
        assert deadline.budgeted(30) == 30
        with deadline.time_budget(500):  # an inner budget can only tighten it
            assert deadline.remaining() <= 60
    assert deadline.remaining() is None


def test_too_little_time_left_means_no_call():
    with deadline.time_budget(deadline.MIN_CALL_SECONDS - 1):
        with pytest.raises(LLMUnavailable, match="Out of time"):
            deadline.budgeted(240)
    with deadline.time_budget(0.01):
        time.sleep(0.03)
        with pytest.raises(LLMUnavailable):
            deadline.check()


def test_the_openai_path_refuses_a_call_it_has_no_time_for():
    client = LLMClient(AIConfig(
        provider="openai", enabled=True, base_url="https://api.example.test/v1",
        api_key="sk-test", model="m",
    ))
    with deadline.time_budget(1):
        with pytest.raises(LLMUnavailable):
            client.complete([{"role": "user", "content": "hi"}])


def test_a_serverless_request_carries_the_budget(monkeypatch):
    from docforge.api.app import app

    seen: dict = {}

    def probe():  # sync, so it runs in the thread pool like the real routes
        seen["left"] = deadline.remaining()
        return {}

    app.add_api_route("/api/_test_budget", probe, methods=["GET"])
    try:
        s = get_settings()
        monkeypatch.setattr(s, "serverless_time_budget_seconds", 240)
        client = TestClient(app)
        monkeypatch.setattr(s, "serverless", True)
        client.get("/api/_test_budget")
        assert seen["left"] is not None and 230 < seen["left"] <= 240
        monkeypatch.setattr(s, "serverless", False)
        client.get("/api/_test_budget")
        assert seen["left"] is None
    finally:
        app.router.routes[:] = [
            r for r in app.router.routes if getattr(r, "path", "") != "/api/_test_budget"
        ]


# --- Running independent calls together -------------------------------------

def test_independent_calls_run_together_and_keep_their_order():
    marker: contextvars.ContextVar[str] = contextvars.ContextVar("marker", default="")
    marker.set("request-7")
    progress_threads = []

    def task(i):
        def run():
            time.sleep(0.3)
            return (i, marker.get(), deadline.remaining() is not None)
        return run

    with deadline.time_budget(100):
        t0 = time.monotonic()
        out = run_all(
            [task(i) for i in range(4)], slots=4,
            on_done=lambda d, n: progress_threads.append(threading.current_thread()),
        )
        elapsed = time.monotonic() - t0

    assert [o[0] for o in out] == [0, 1, 2, 3]
    assert all(o[1] == "request-7" and o[2] for o in out)  # context travels along
    assert elapsed < 0.9, elapsed  # 4 x 0.3 s one after another would be 1.2 s
    assert progress_threads == [threading.main_thread()] * 4


def test_a_cancellation_wins_over_other_errors():
    def fail(exc):
        def run():
            raise exc
        return run

    with pytest.raises(LLMCancelled):
        run_all([fail(LLMError("a")), fail(LLMCancelled("b"))], slots=2)
    with pytest.raises(LLMError, match="first"):
        run_all([lambda: 1, fail(LLMError("first")), fail(LLMError("second"))], slots=3)


def test_local_model_servers_get_one_call_at_a_time():
    assert slots_for("http://localhost:11434/v1") == 1
    assert slots_for("http://host.docker.internal:1234/v1") == 1
    assert slots_for("https://api.anthropic.com") == get_settings().ai_max_parallel_calls > 1


# --- End to end: an analysis that used to run out of time --------------------

def _long_docx(tmp_path, n=90):
    doc = Document()
    for i in range(n):
        if i % 10 == 0:
            doc.add_heading(f"Section {i // 10 + 1}", level=1)
        doc.add_paragraph(f"Paragraph {i} explains one part of the plan in a full sentence.")
    path = tmp_path / "long.docx"
    doc.save(str(path))
    return path


def _ai_client():
    return LLMClient(AIConfig(
        provider="anthropic", enabled=True, base_url="https://api.anthropic.com",
        api_key="sk-ant-test", model="claude-haiku-5-5",
    ))


def test_classify_batches_run_together(tmp_path, monkeypatch):
    from docforge.ai_classifier import classify
    from docforge.structure_normalizer import build_extraction

    ext = build_extraction(_long_docx(tmp_path), "d0")
    n_top = len(ext.top_level_elements())
    assert n_top > 80  # at least five classify batches of 20
    calls: list[str] = []

    def fake_agentic(self, *, schema, **kw):
        calls.append(schema.__name__)
        time.sleep(0.3)
        return schema()

    monkeypatch.setattr(LLMClient, "complete_agentic", fake_agentic)
    t0 = time.monotonic()
    result = classify(ext, None, client=_ai_client(), mode="smart")
    elapsed = time.monotonic() - t0

    batches = calls.count("LLMClassifyResponse")
    assert batches >= 5
    # understand + critique + one round of batches, instead of every batch in turn
    assert elapsed < 0.3 * (batches + 2) * 0.7, (elapsed, batches)
    assert result is not None


def test_an_analysis_out_of_time_still_finishes(tmp_path, monkeypatch):
    from docforge.ai_classifier import classify
    from docforge.structure_normalizer import build_extraction

    ext = build_extraction(_long_docx(tmp_path), "d0")

    def slow_agentic(self, *, schema, **kw):
        deadline.budgeted(self.config.timeout_seconds)  # what every transport does
        time.sleep(0.2)
        return schema()

    monkeypatch.setattr(LLMClient, "complete_agentic", slow_agentic)
    monkeypatch.setattr(deadline, "MIN_CALL_SECONDS", 0.15)
    with deadline.time_budget(0.25):  # room for about one call
        result = classify(ext, None, client=_ai_client(), mode="tags_only")

    assert result.source == "heuristic_fallback"
    assert "Out of time" in (result.ai_warning or "")
    assert result.classifications  # a complete template, from the built-in engine
