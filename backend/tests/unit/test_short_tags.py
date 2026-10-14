"""Short, readable tags: short field names, and no {%p if %} wrapper per field.

A template is read by people — in the "with tags" view and in Word. Every tag
used to be up to five words of the example text plus its section number, and
every field paragraph sat between its own {%p if %} and {%p endif %} lines.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from docforge.ai_classifier import classify, derive_field_definitions
from docforge.ai_classifier.fields import include_field_name
from docforge.assembler import assemble
from docforge.common.textutil import short_field_name
from docforge.schemas.classification import ElementClassification
from docforge.schemas.enums import ClassificationType, FieldType
from docforge.schemas.template import FieldDefinition
from docforge.structure_normalizer import build_extraction
from docforge.template_builder import build_template_docx


@pytest.mark.parametrize(
    "text, expected",
    [
        ("4.2 SLA Breach Analysis", "sla_breach_analysis"),
        ("1. Executive Summary", "executive_summary"),
        ("Project Name:", "project_name"),
        ("The following critical findings were identified", "following_critical"),
        ("Data Protection and Privacy", "data_protection_privacy"),
        ("2025", "f_2025"),
        ("", "field"),
    ],
)
def test_short_field_name(text, expected):
    assert short_field_name(text) == expected


@pytest.mark.parametrize(
    "text",
    ["4.2 SLA Breach Analysis", "Overall Compliance Rating", "key_findings_2", "x" * 60, "2025"],
)
def test_short_field_name_is_short_and_idempotent(text):
    name = short_field_name(text)
    assert len(name) <= 24 and name.count("_") <= 3
    assert short_field_name(name) == name


def test_include_toggle_never_says_include_twice():
    c = ElementClassification(
        node_id="n1", field_name="include_overall_vendor_achieved_compliance", optional=True
    )
    assert include_field_name(c, None) == "include_overall_vendor"


def _numbered_report(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Acme Financial Group"
    doc.add_heading("4.2 Service Level Agreement Performance Review", level=1)
    doc.add_paragraph("Uptime was 99.95% across the quarter.")
    doc.add_heading("4.3 Equipment", level=1)
    t = doc.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text = "Item"
    t.rows[0].cells[1].text = "Status"
    t.rows[1].cells[0].text = "Projector"
    t.rows[1].cells[1].text = "Working"
    t.rows[2].cells[0].text = "Screen"
    t.rows[2].cells[1].text = "Broken"
    doc.add_paragraph("All equipment was inspected on site.")
    path = tmp_path / "report.docx"
    doc.save(str(path))
    return path


def test_tags_only_names_are_short_and_carry_no_section_numbers(tmp_path):
    path = _numbered_report(tmp_path)
    ext = build_extraction(path, "d0")
    fields = derive_field_definitions(ext, classify(ext, None, mode="tags_only"))
    names = [f.field_name for f in fields]
    assert names, "tags-only mode fields every text"
    for name in names:
        assert not name.startswith("f_4"), name
        assert len(name) <= 30, name
    # A column heading is named after itself, not after its section.
    assert "status" in names
    # Page header text is named for the header, not the last body section.
    assert any(n.startswith("header") for n in names), names


def test_tags_only_template_has_one_line_per_field(tmp_path):
    path = _numbered_report(tmp_path)
    ext = build_extraction(path, "d0")
    result = classify(ext, None, mode="tags_only")
    fields = derive_field_definitions(ext, result)
    built = Document(BytesIO(build_template_docx(path, result, fields)))
    text = "\n".join(p.text for p in built.paragraphs)
    for f in fields:
        assert f"{{%p if {f.field_name} %}}" not in text, f"{f.field_name} is still wrapped"
    assert "is not defined" not in text


# --------------------------------------------------------------------------- #
# Render-time rules that replaced the tags                                     #
# --------------------------------------------------------------------------- #
def _template(build) -> bytes:
    doc = Document()
    build(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _texts(data: bytes) -> list[str]:
    return [p.text for p in Document(BytesIO(data)).paragraphs]


def _field(name: str) -> FieldDefinition:
    return FieldDefinition(field_name=name, field_type=FieldType.TEXT)


def test_an_empty_lone_placeholder_takes_its_paragraph_with_it():
    tpl = _template(lambda d: [d.add_paragraph(t) for t in ("Intro", "{{ notes }}", "End")])
    fields = [_field("notes")]
    assert _texts(assemble(tpl, {"notes": ""}, fields)) == ["Intro", "End"]
    assert _texts(assemble(tpl, {"notes": "Hello"}, fields)) == ["Intro", "Hello", "End"]


def test_a_placeholder_beside_text_keeps_its_paragraph():
    tpl = _template(lambda d: d.add_paragraph("Owner: {{ owner }}"))
    assert _texts(assemble(tpl, {"owner": ""}, [_field("owner")])) == ["Owner: "]


def test_a_name_the_caller_never_supplied_is_left_alone():
    # A loop variable or a field removed in review: render as before (blank).
    tpl = _template(lambda d: [d.add_paragraph(t) for t in ("A", "{{ ghost }}", "B")])
    assert _texts(assemble(tpl, {}, None)) == ["A", "", "B"]


def test_a_table_cell_keeps_its_only_paragraph():
    def build(d):
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "Label"
        t.cell(0, 1).text = "{{ value }}"

    out = Document(BytesIO(assemble(_template(build), {"value": ""}, [_field("value")])))
    cell = out.tables[0].cell(0, 1)
    assert len(cell.paragraphs) == 1 and cell.text == ""


def test_section_toggles_default_to_shown():
    def build(d):
        d.add_paragraph("{%p if _show_risks %}")
        d.add_paragraph("Risk text")
        d.add_paragraph("{%p endif %}")

    tpl = _template(build)
    assert _texts(assemble(tpl, {}, None)) == ["Risk text"]
    assert _texts(assemble(tpl, {"_show_risks": False}, None)) == []


def test_a_fully_filled_template_is_rendered_from_the_same_bytes():
    from docforge.assembler.placeholders import prepare_template

    tpl = _template(lambda d: d.add_paragraph("{{ notes }}"))
    assert prepare_template(tpl, {"notes": "x"}) is tpl


def test_a_tags_only_heading_classification_still_renders(tmp_path):
    # The classic case the old wrapper existed for: an unfilled heading field.
    path = _numbered_report(tmp_path)
    ext = build_extraction(path, "d0")
    result = classify(ext, None, mode="tags_only")
    fields = derive_field_definitions(ext, result)
    tpl = build_template_docx(path, result, fields)
    titles = [f for f in fields if f.field_name.endswith("_title")]
    assert titles
    ctx = {f.field_name: "" for f in fields}
    ctx[titles[0].field_name] = "Service levels"
    out = _texts(assemble(tpl, ctx, fields))
    assert "Service levels" in out
    assert "" not in out
    assert all(
        c.classification != ClassificationType.FIXED or not c.field_name
        for c in result.classifications
    )
